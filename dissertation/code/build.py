"""
build.py - assemble the dissertation (Markdown sources in src/<lang>/) into
build/<lang>/Dissertation_<lang>.docx and .pdf.

Usage: python3 build.py en|ru|uz [--passes 2]

Pipeline: preprocess placeholders (figures, tables, equations, cross-references,
citations, counts) -> pandoc (reference.docx with the formatting rules) ->
python-docx post-processing (title page, table borders, numbered equations,
static table of contents, page numbers) -> LibreOffice PDF -> page numbers
of headings are read from the PDF and the contents are rebuilt (2 passes).
"""
import copy, glob, json, os, re, subprocess, sys
import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
D = os.environ.get('DISS_ROOT', os.path.join(HERE, '..'))
REFS = json.load(open(os.path.join(D, 'data', 'refs.json')))

CFG = {
    'en': dict(fig='Fig. {}', figcap='Fig. {}. ', tab='Table {}', tabcap='Table {}. ', eq='({})',
               chap=r'^#\s+CHAPTER\s+(\d+)\.', toc='CONTENTS', refs='REFERENCES', ann='ANNOTATION',
               ann_titles={'en': 'ANNOTATION', 'ru': 'АННОТАЦИЯ', 'uz': 'ANNOTATSIYA'}),
    'ru': dict(fig='рис. {}', figcap='Рис. {}. ', tab='табл. {}', tabcap='Таблица {}. ', eq='({})',
               chap=r'^#\s+ГЛАВА\s+(\d+)\.', toc='ОГЛАВЛЕНИЕ', refs='СПИСОК ИСПОЛЬЗОВАННОЙ ЛИТЕРАТУРЫ', ann='АННОТАЦИЯ',
               ann_titles={'en': 'ANNOTATION', 'ru': 'АННОТАЦИЯ', 'uz': 'ANNOTATSIYA'}),
    'uz': dict(fig='{}-rasm', figcap='{}-rasm. ', tab='{}-jadval', tabcap='{}-jadval. ', eq='({})',
               chap=r'^#\s+([IVX]+)\s+BOB\.', toc='MUNDARIJA', refs='FOYDALANILGAN ADABIYOTLAR RO‘YXATI', ann='ANNOTATSIYA',
               ann_titles={'en': 'ANNOTATION', 'ru': 'АННОТАЦИЯ', 'uz': 'ANNOTATSIYA'}),
}
ROMAN = {'I': 1, 'II': 2, 'III': 3, 'IV': 4, 'V': 5}

TITLE = {
    'en': dict(ministry='MINISTRY OF HIGHER EDUCATION, SCIENCE AND INNOVATION OF THE REPUBLIC OF UZBEKISTAN',
               univ='[NAME OF THE UNIVERSITY]', fac='Faculty: [name of the faculty]', dep='Department: [name of the department]',
               rights='As a manuscript', udc='UDC 539.12.04:615.849', author='[FULL NAME OF THE STUDENT]',
               title='X-RAY SPECTRUM MODELING FOR THE STUDY OF RADIO-ENHANCEMENT IN BIOLOGICAL MATTER WITH Ag',
               spec='Specialty: [code] – [name of the master’s specialty]',
               kind='MASTER’S DISSERTATION\nfor the academic degree of Master',
               sup='Scientific supervisor:\n[academic degree, title, full name]',
               admit='“Admitted to defence”\nHead of the department ____________ [full name]\n“___” ____________ 2026',
               city='[City] – 2026'),
    'ru': dict(ministry='МИНИСТЕРСТВО ВЫСШЕГО ОБРАЗОВАНИЯ, НАУКИ И ИННОВАЦИЙ РЕСПУБЛИКИ УЗБЕКИСТАН',
               univ='[НАЗВАНИЕ УНИВЕРСИТЕТА]', fac='Факультет: [название факультета]', dep='Кафедра: [название кафедры]',
               rights='На правах рукописи', udc='УДК 539.12.04:615.849', author='[ФАМИЛИЯ ИМЯ ОТЧЕСТВО СТУДЕНТА]',
               title='МОДЕЛИРОВАНИЕ СПЕКТРА РЕНТГЕНОВСКОГО ИЗЛУЧЕНИЯ ДЛЯ ИССЛЕДОВАНИЯ РАДИОУСИЛЕНИЯ В БИОЛОГИЧЕСКОЙ СРЕДЕ С Ag',
               spec='Специальность: [шифр] – [название магистерской специальности]',
               kind='МАГИСТЕРСКАЯ ДИССЕРТАЦИЯ\nна соискание академической степени магистра',
               sup='Научный руководитель:\n[учёная степень, звание, Ф. И. О.]',
               admit='«Допускается к защите»\nЗаведующий кафедрой ____________ [Ф. И. О.]\n«___» ____________ 2026 г.',
               city='[Город] – 2026'),
    'uz': dict(ministry='O‘ZBEKISTON RESPUBLIKASI OLIY TA’LIM, FAN VA INNOVATSIYALAR VAZIRLIGI',
               univ='[UNIVERSITET NOMI]', fac='Fakultet: [fakultet nomi]', dep='Kafedra: [kafedra nomi]',
               rights='Qo‘lyozma huquqida', udc='UO‘K 539.12.04:615.849', author='[TALABANING FAMILIYASI, ISMI, OTASINING ISMI]',
               title='BIOLOGIK MUHITDA Ag BILAN RADIOKUCHAYISHNI TADQIQ QILISH UCHUN RENTGEN NURLANISHI SPEKTRINI MODELLASHTIRISH',
               spec='Mutaxassislik: [shifr] – [magistratura mutaxassisligi nomi]',
               kind='Magistr akademik darajasini olish uchun yozilgan\nMAGISTRLIK DISSERTATSIYASI',
               sup='Ilmiy rahbar:\n[ilmiy darajasi, unvoni, F. I. Sh.]',
               admit='“Himoyaga ruxsat etiladi”\nKafedra mudiri ____________ [F. I. Sh.]\n“___” ____________ 2026-yil',
               city='[Shahar] – 2026'),
}


# ----------------------------------------------------------------- sources
def read_sources(lang):
    files = sorted(glob.glob(os.path.join(D, 'src', lang, '*.md')))
    parts = {}
    for f in files:
        parts[os.path.basename(f)] = open(f, encoding='utf-8').read().strip() + '\n'
    return parts


def body_text(parts):
    keys = [k for k in sorted(parts) if not k.startswith('00_')]
    return '\n\n'.join(parts[k] for k in keys)


def citation_order():
    """reference numbering by first appearance in the English text"""
    en = read_sources('en')
    txt = parts_ordered_text(en)
    order = []
    for m in re.finditer(r'\[(@[^\]]+)\]', txt):
        for k in re.findall(r'@([A-Za-z0-9_]+)', m.group(1)):
            if k not in order:
                order.append(k)
    return order


def parts_ordered_text(parts):
    keys = [k for k in sorted(parts) if k not in ('00_annotation.md',)]
    return '\n\n'.join(parts[k] for k in keys)


def fmt_cite(nums):
    nums = sorted(set(nums))
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            out.append(f'{nums[i]}–{nums[j]}')
        else:
            out.extend(str(n) for n in nums[i:j + 1])
        i = j + 1
    return '[' + ', '.join(out) + ']'


# ----------------------------------------------------------------- preprocessing
def preprocess(lang, text, order, counts):
    c = CFG[lang]
    labels, figs, tabs = {}, [], []
    chap = 0
    nfig = ntab = neq = 0
    out = []
    lines = text.split('\n')
    i = 0
    # first pass: numbering (in reading order)
    for ln in lines:
        m = re.match(c['chap'], ln)
        if m:
            chap = ROMAN.get(m.group(1), None) or int(m.group(1))
            nfig = ntab = neq = 0
        elif ln.startswith('# '):
            chap = 0 if not ln.startswith('# APPEND') and 'ILOVA' not in ln and 'ПРИЛОЖ' not in ln else chap
        for fm in re.finditer(r'\{\{FIG:([\w]+)\|', ln):
            nfig += 1
            labels[fm.group(1)] = c['fig'].format(f'{chap}.{nfig}')
            labels['__cap_' + fm.group(1)] = c['figcap'].format(f'{chap}.{nfig}')
            figs.append(fm.group(1))
        for tm in re.finditer(r'\{\{TABLE:([\w]+)\}\}', ln):
            ntab += 1
            labels[tm.group(1)] = c['tab'].format(f'{chap}.{ntab}')
            labels['__cap_' + tm.group(1)] = c['tabcap'].format(f'{chap}.{ntab}')
            tabs.append(tm.group(1))
        tm = re.match(r'^Table:\s*\{#(tab:[\w]+)\}', ln)
        if tm and tm.group(1) != 'tab:abbrev':
            ntab += 1
            labels[tm.group(1)] = c['tab'].format(f'{chap}.{ntab}')
            labels['__cap_' + tm.group(1)] = c['tabcap'].format(f'{chap}.{ntab}')
            tabs.append(tm.group(1))
        em = re.search(r'\$\$\s*\{#(eq:[\w]+)\}', ln)
        if em:
            neq += 1
            labels[em.group(1)] = c['eq'].format(f'{chap}.{neq}')
    # second pass: substitutions
    text = '\n'.join(lines)

    def fig_sub(m):
        fid, cap = m.group(1), m.group(2).strip()
        path = os.path.join(D, 'figures', lang, fid + '.png')
        return f'\n\n![{labels["__cap_" + fid]}{cap}]({path}){{width=13.5cm}}\n\n'
    text = re.sub(r'\{\{FIG:([\w]+)\|(.+?)\}\}', fig_sub, text, flags=re.S)

    def tab_sub(m):
        tid = m.group(1)
        raw = open(os.path.join(D, 'data', 'tables', lang, tid + '.md'), encoding='utf-8').read().strip().split('\n')
        title, body = raw[0], '\n'.join(raw[2:])
        return f'\n\nTable: {labels["__cap_" + tid]}{title}\n\n{body}\n\n'
    text = re.sub(r'\{\{TABLE:([\w]+)\}\}', tab_sub, text)
    text = re.sub(r'^Table:\s*\{#tab:abbrev\}.*$', '', text, flags=re.M)
    text = re.sub(r'^Table:\s*\{#(tab:[\w]+)\}\s*',
                  lambda m: 'Table: ' + labels.get('__cap_' + m.group(1), ''), text, flags=re.M)
    # equations: $$ ... $$ {#eq:x}  ->  $$...$$ + marker paragraph
    def eq_sub(m):
        body, lab = m.group(1).strip(), m.group(2)
        return f'$${body}$$\n\nEQNUM[{labels[lab][1:-1]}]\n'
    text = re.sub(r'\$\$(.+?)\$\$\s*\{#(eq:[\w]+)\}', eq_sub, text, flags=re.S)
    # cross references
    def ref_sub(m):
        k = m.group(1)
        if k not in labels:
            print('WARNING unresolved ref', k)
            return '??'
        return labels[k]
    text = re.sub(r'\{\{REF:([\w:]+)\}\}', ref_sub, text)
    # equation references are already parenthesised: "((2.4))" -> "(2.4)"
    text = re.sub(r'\(\((\d+\.\d+)\)\)', r'(\1)', text)
    # citations
    def cite_sub(m):
        keys = re.findall(r'@([A-Za-z0-9_]+)', m.group(1))
        nums = []
        for k in keys:
            if k not in order:
                print('WARNING citation not in EN order:', k)
                order.append(k)
            nums.append(order.index(k) + 1)
        return fmt_cite(nums)
    text = re.sub(r'\[(@[A-Za-z0-9_]+(?:\s*;\s*@[A-Za-z0-9_]+)*)\]', cite_sub, text)
    # counts
    for k, v in counts.items():
        text = text.replace('{{COUNT:%s}}' % k, str(v))
    return text, labels, figs, tabs


def references_md(lang, order):
    c = CFG[lang]
    lines = [f'# {c["refs"]}', '']
    for i, k in enumerate(order, 1):
        if k not in REFS:
            print('WARNING unknown reference', k)
            continue
        t = REFS[k]['text'].replace('*', '\\*')
        lines.append(f'::: {{custom-style="Reference"}}\n{i}. {t}\n:::\n')
    return '\n'.join(lines)


# ----------------------------------------------------------------- reference.docx
def make_reference(path):
    tmp = path + '.tmp.docx'
    open(tmp, 'wb').write(subprocess.run(['pandoc', '--print-default-data-file', 'reference.docx'],
                                         capture_output=True, check=True).stdout)
    doc = docx.Document(tmp)
    from docx.enum.style import WD_STYLE_TYPE
    _styles = doc.styles

    class _St:
        def __getitem__(self, name):
            for s_ in _styles:
                if s_.name == name:
                    return s_
            s_ = _styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s_.base_style = [x for x in _styles if x.name == 'Normal'][0]
            return s_

        def __iter__(self):
            return iter(_styles)

        def add_style(self, *a):
            return _styles.add_style(*a)
    st = _St()

    def font(s, size=14, bold=None, italic=False, name='Times New Roman', color=True):
        s.font.name = name
        s.font.size = Pt(size)
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a), name)
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        if bold is not None:
            s.font.bold = bold
        s.font.italic = italic
        if color:
            s.font.color.rgb = RGBColor(0, 0, 0)

    def para(s, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=1.25, spacing=1.5, before=0, after=0, left=0, keep=False):
        pf = s.paragraph_format
        pf.alignment = align
        pf.first_line_indent = Cm(first) if first else Cm(0)
        pf.left_indent = Cm(left)
        pf.line_spacing = spacing
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = keep
        pf.widow_control = True

    for name in ['Normal', 'Body Text', 'First Paragraph']:
        font(st[name], 14)
        para(st[name])
    # 'Compact' is used by pandoc for tight list items and table cells; cells are reformatted later
    font(st['Compact'], 14)
    para(st['Compact'], align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=0, spacing=1.5)
    font(st['Heading 1'], 14, bold=True)
    para(st['Heading 1'], align=WD_ALIGN_PARAGRAPH.CENTER, first=0, spacing=1.5, before=0, after=12, keep=True)
    st['Heading 1'].paragraph_format.page_break_before = True
    font(st['Heading 2'], 14, bold=True)
    para(st['Heading 2'], align=WD_ALIGN_PARAGRAPH.LEFT, first=1.25, spacing=1.5, before=12, after=6, keep=True)
    font(st['Heading 3'], 14, bold=True, italic=True)
    para(st['Heading 3'], align=WD_ALIGN_PARAGRAPH.LEFT, first=1.25, spacing=1.5, before=6, after=3, keep=True)
    for name in ['Image Caption', 'Caption']:
        font(st[name], 12)
        para(st[name], align=WD_ALIGN_PARAGRAPH.CENTER, first=0, spacing=1.0, before=2, after=10)
    font(st['Table Caption'], 12)
    para(st['Table Caption'], align=WD_ALIGN_PARAGRAPH.LEFT, first=0, spacing=1.0, before=8, after=4, keep=True)
    for name in ['Figure', 'Captioned Figure']:
        para(st[name], align=WD_ALIGN_PARAGRAPH.CENTER, first=0, spacing=1.0, before=6, after=0, keep=True)
    font(st['Source Code'], 8, name='Courier New')
    para(st['Source Code'], align=WD_ALIGN_PARAGRAPH.LEFT, first=0, spacing=1.0)
    # custom styles
    for nm, kw in [('Reference', dict(size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=0, spacing=1.0, after=2, left=0)),
                   ('TitleCenter', dict(size=14, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, spacing=1.15, after=0)),
                   ('TitleRight', dict(size=14, align=WD_ALIGN_PARAGRAPH.RIGHT, first=0, spacing=1.15, after=0)),
                   ('TitleLeft', dict(size=14, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, spacing=1.15, after=0)),
                   ('TitleBig', dict(size=16, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, spacing=1.15, after=0)),
                   ('Keywords', dict(size=14, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=1.25, spacing=1.5, after=0)),
                   ('Annotation', dict(size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=1.0, spacing=1.15, after=3)),
                   ('TOCLine1', dict(size=14, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, spacing=1.15, after=2)),
                   ('TOCLine2', dict(size=14, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, spacing=1.15, after=2, left=0.75))]:
        if nm not in [s.name for s in st]:
            s = st.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = st['Normal']
        s = st[nm]
        font(s, kw['size'], bold=True if nm == 'TitleBig' else None)
        para(s, align=kw['align'], first=kw['first'], spacing=kw['spacing'], after=kw.get('after', 0), left=kw.get('left', 0))
    if 'Reference' in [s.name for s in st]:
        st['Reference'].paragraph_format.left_indent = Cm(0.75)
        st['Reference'].paragraph_format.first_line_indent = Cm(-0.75)
    # page setup
    for sec in doc.sections:
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        sec.footer_distance = Cm(1.0)
    doc.save(path)
    os.remove(tmp)


# ----------------------------------------------------------------- post-processing
TEXT_W = 16.5   # cm


def S(doc, name):
    for s_ in doc.styles:
        if s_.name == name:
            return s_
    raise KeyError(name)


def set_cell_borders(table, val='single', sz=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement(f'w:{e}')
        b.set(qn('w:val'), val)
        b.set(qn('w:sz'), str(sz))
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '000000')
        borders.append(b)
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(borders)


def set_table_width(table, cm):
    tblPr = table._tbl.tblPr
    w = tblPr.find(qn('w:tblW'))
    if w is None:
        w = OxmlElement('w:tblW')
        tblPr.append(w)
    w.set(qn('w:type'), 'dxa')
    w.set(qn('w:w'), str(int(cm / 2.54 * 1440)))
    lay = tblPr.find(qn('w:tblLayout'))
    if lay is None:
        lay = OxmlElement('w:tblLayout')
        tblPr.append(lay)
    lay.set(qn('w:type'), 'autofit')


def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for el in list(p._p):
        if el.tag != qn('w:pPr'):
            p._p.remove(el)
    run = p.add_run()
    for kind, txt in [('begin', None), (None, ' PAGE '), ('separate', None), (None, '1'), ('end', None)]:
        if kind:
            fc = OxmlElement('w:fldChar')
            fc.set(qn('w:fldCharType'), kind)
            run._r.append(fc)
        elif txt.strip() == 'PAGE':
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = txt
            run._r.append(it)
        else:
            t = OxmlElement('w:t')
            t.text = txt
            run._r.append(t)
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'


def number_equations(doc):
    body = doc.element.body
    paras = list(body.iterchildren(qn('w:p')))
    n = 0
    for i, p in enumerate(paras):
        txt = ''.join(t.text or '' for t in p.iter(qn('w:t')))
        m = re.fullmatch(r'\s*EQNUM\[([\d.]+)\]\s*', txt)
        if not m:
            continue
        # previous sibling paragraph with the display equation
        prev = p.getprevious()
        while prev is not None and prev.tag != qn('w:p'):
            prev = prev.getprevious()
        mp = prev.find('.//' + qn('m:oMathPara')) if prev is not None else None
        if mp is None:
            print('WARNING equation marker without math', m.group(1))
            body.remove(p)
            continue
        # build a borderless 2-cell table: equation | (number)
        tbl = doc.add_table(rows=1, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_cell_borders(tbl, 'nil', 0)
        c0, c1 = tbl.rows[0].cells
        c0.width, c1.width = Cm(TEXT_W - 1.8), Cm(1.8)
        for c in (c0, c1):
            c.vertical_alignment = 1  # center
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.first_line_indent = Cm(0)
        p0.paragraph_format.space_before = Pt(4)
        p0.paragraph_format.space_after = Pt(4)
        p0._p.append(copy.deepcopy(mp))
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.first_line_indent = Cm(0)
        r = p1.add_run(f'({m.group(1)})')
        r.font.size = Pt(14)
        r.font.name = 'Times New Roman'
        prev.addnext(tbl._tbl)
        body.remove(prev)
        body.remove(p)
        n += 1
    return n


def format_tables(doc):
    for t in doc.tables:
        tblPr = t._tbl.tblPr
        style = tblPr.find(qn('w:tblStyle'))
        borders = tblPr.find(qn('w:tblBorders'))
        if borders is not None and borders.find(qn('w:top')) is not None and borders.find(qn('w:top')).get(qn('w:val')) == 'nil':
            continue      # equation tables
        set_cell_borders(t)
        set_table_width(t, TEXT_W)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(t.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.line_spacing = 1.0
                    if p.paragraph_format.alignment is None or p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
                        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.size = Pt(10)
                        if ri == 0:
                            r.font.bold = True
            if ri == 0:
                trPr = row._tr.get_or_add_trPr()
                h = OxmlElement('w:tblHeader')
                trPr.append(h)


def insert_title_page(doc, lang):
    T = TITLE[lang]
    body = doc.element.body
    first = body[0]
    paras = []

    def P(text, style, bold=False, space_after=0, size=None, caps=False):
        p = doc.add_paragraph(style=S(doc, style))
        lines = text.split('\n')
        for i, ln in enumerate(lines):
            r = p.add_run(ln)
            r.bold = bold
            if size:
                r.font.size = Pt(size)
            if i < len(lines) - 1:
                r.add_break()
        p.paragraph_format.space_after = Pt(space_after)
        paras.append(p)
        return p

    P(T['ministry'], 'TitleCenter', bold=True, space_after=6)
    P(T['univ'], 'TitleCenter', bold=True, space_after=6)
    P(T['fac'] + '\n' + T['dep'], 'TitleCenter', space_after=30)
    P(T['udc'], 'TitleLeft', space_after=0)
    P(T['rights'], 'TitleRight', space_after=40)
    P(T['author'], 'TitleCenter', bold=True, space_after=30)
    P(T['title'], 'TitleBig', bold=True, space_after=30)
    P(T['spec'], 'TitleCenter', space_after=24)
    P(T['kind'], 'TitleCenter', bold=True, space_after=50)
    P(T['admit'], 'TitleLeft', space_after=24, size=13)
    P(T['sup'], 'TitleRight', space_after=60)
    last = P(T['city'], 'TitleCenter', bold=True)
    last.add_run().add_break(WD_BREAK.PAGE)
    for p in paras:
        first.addprevious(p._p)


def toc_entries(doc):
    out = []
    for p in doc.paragraphs:
        sn = p.style.name
        if sn in ('Heading 1', 'Heading 2') and p.text.strip():
            out.append((1 if sn == 'Heading 1' else 2, p.text.strip()))
    return out


def insert_toc(doc, lang, entries, pages):
    """static contents after the annotation pages (before the abbreviations heading)"""
    c = CFG[lang]
    body = doc.element.body
    # insert before the first Heading 1 that is not an annotation
    anchor = None
    for p in doc.paragraphs:
        if p.style.name == 'Heading 1' and p.text.strip() not in c['ann_titles'].values():
            anchor = p
            break
    head = doc.add_paragraph(c['toc'], style=S(doc, 'Heading 1'))
    items = [head]
    for lvl, txt in entries:
        if txt in c['ann_titles'].values():
            continue
        p = doc.add_paragraph(style=S(doc, 'TOCLine1' if lvl == 1 else 'TOCLine2'))
        p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if lvl == 1:
            p.paragraph_format.left_indent = Cm(0)
            p.paragraph_format.first_line_indent = Cm(0)
        else:
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.first_line_indent = Cm(-0.5)
        short = txt
        r = p.add_run(short)
        if lvl == 1:
            r.bold = True
        p.add_run('\t' + str(pages.get(txt, '')))
        items.append(p)
    for it in items:
        anchor._p.addprevious(it._p)


def fix_math(doc):
    """explicit delimiters (LibreOffice defaults differ from Word) and centred display math"""
    M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    body = doc.element.body
    for d in body.iter('{%s}d' % M):
        dpr = d.find('{%s}dPr' % M)
        if dpr is None:
            dpr = OxmlElement('m:dPr')
            d.insert(0, dpr)
        if dpr.find('{%s}begChr' % M) is None:
            b = OxmlElement('m:begChr'); b.set(qn('m:val'), '('); dpr.insert(0, b)
        if dpr.find('{%s}endChr' % M) is None:
            e = OxmlElement('m:endChr'); e.set(qn('m:val'), ')'); dpr.append(e)
    for mp in body.iter('{%s}oMathPara' % M):
        ppr = mp.find('{%s}oMathParaPr' % M)
        if ppr is None:
            ppr = OxmlElement('m:oMathParaPr')
            mp.insert(0, ppr)
        jc = ppr.find('{%s}jc' % M)
        if jc is None:
            jc = OxmlElement('m:jc')
            ppr.append(jc)
        jc.set(qn('m:val'), 'center')


def postprocess(path, lang, pages):
    doc = docx.Document(path)
    fix_math(doc)
    neq = number_equations(doc)
    format_tables(doc)
    entries = toc_entries(doc)
    insert_toc(doc, lang, entries, pages)
    insert_title_page(doc, lang)
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    add_page_number_footer(sec)
    # language of the text (spell checking)
    langcode = {'en': 'en-GB', 'ru': 'ru-RU', 'uz': 'uz-Latn-UZ'}[lang]
    for s in doc.styles:
        try:
            rpr = s.element.get_or_add_rPr()
            l_ = rpr.find(qn('w:lang'))
            if l_ is None:
                l_ = OxmlElement('w:lang')
                rpr.append(l_)
            l_.set(qn('w:val'), langcode)
        except Exception:
            pass
    doc.core_properties.title = TITLE[lang]['title']
    doc.core_properties.author = ''
    doc.save(path)
    return entries, neq


# ----------------------------------------------------------------- pdf + page map
def to_pdf(docx_path):
    outdir = os.path.dirname(docx_path)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', outdir, docx_path],
                   check=True, capture_output=True, timeout=600)
    return docx_path[:-5] + '.pdf'


def norm(s):
    s = s.replace(' ', ' ').replace('’', "'").replace('‘', "'").replace('ʻ', "'").replace('ʼ', "'")
    return re.sub(r'[\s\-–—.,:;()]+', '', s).lower()


def page_map(pdf, entries, lang):
    n = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout).group(1))
    raw = [subprocess.run(['pdftotext', '-f', str(i), '-l', str(i), '-layout', pdf, '-'], capture_output=True, text=True).stdout
           for i in range(1, n + 1)]
    texts = [norm(t) for t in raw]
    c = CFG[lang]
    # the contents occupy the pages with many dot leaders, starting at the page titled CONTENTS
    toc_first = next((i for i, t in enumerate(raw) if c['toc'] in t), 0)
    start = toc_first + 1
    while start < n and len(re.findall(r'\.{6,}', raw[start])) >= 5:
        start += 1
    pages, cur = {}, start
    for lvl, txt in entries:
        if txt in c['ann_titles'].values():
            continue
        key = norm(txt)[:60]
        for j in range(cur, n):
            if key and key in texts[j]:
                pages[txt] = j + 1
                cur = j
                break
    return pages, n


# ----------------------------------------------------------------- main
def build(lang, passes=2):
    out = os.path.join(D, 'build', lang)
    os.makedirs(out, exist_ok=True)
    ref = os.path.join(D, 'build', 'reference.docx')
    if not os.path.exists(ref):
        make_reference(ref)
    parts = read_sources(lang)
    order = citation_order()
    c = CFG[lang]
    # annotations in all three languages (this language first)
    ann = []
    for L in [lang] + [x for x in ('uz', 'ru', 'en') if x != lang]:
        f = os.path.join(D, 'src', L, '00_annotation.md')
        if os.path.exists(f):
            a = open(f, encoding='utf-8').read().strip()
            a = re.sub(r'^#\s+.*$', '', a, count=1, flags=re.M).strip()
            ann.append('# ' + c['ann_titles'][L] + '\n\n::: {custom-style="Annotation"}\n' + a + '\n:::')
    abbrev = parts.get('00_abbrev.md', '')
    main = parts_ordered_text({k: v for k, v in parts.items() if not k.startswith('00_')})
    counts = {'pages': 90, 'figures': 0, 'tables': 0, 'refs': len(order)}
    pages = {}
    for it in range(passes):
        md_main, labels, figs, tabs = preprocess(lang, abbrev + '\n\n' + main, order, counts)
        counts['figures'], counts['tables'] = len(figs), len(tabs)
        md = '\n\n'.join(ann) + '\n\n' + md_main + '\n\n' + references_md(lang, order)
        # move appendices after the references
        m = re.search(r'^# (APPENDICES|ПРИЛОЖЕНИЯ|ILOVALAR)\b.*$', md, flags=re.M)
        if m:
            app_start = m.start()
            refs_start = md.index('# ' + c['refs'])
            if app_start < refs_start:
                app = md[app_start:refs_start]
                md = md[:app_start] + md[refs_start:] + '\n\n' + app
        md_path = os.path.join(out, f'dissertation_{lang}.md')
        open(md_path, 'w', encoding='utf-8').write(md)
        dx = os.path.join(out, f'Dissertation_{lang}.docx')
        subprocess.run(['pandoc', md_path, '-f', 'markdown+tex_math_dollars-auto_identifiers', '-o', dx,
                        '--reference-doc', ref, '--resource-path', D], check=True)
        entries, neq = postprocess(dx, lang, pages)
        pdf = to_pdf(dx)
        pages, n = page_map(pdf, entries, lang)
        counts['pages'] = n
        print(f'[{lang}] pass {it + 1}: {n} pages, {len(figs)} figures, {len(tabs)} tables, {neq} equations, '
              f'{len(order)} references, {len(pages)}/{len(entries)} headings located')
    return n


if __name__ == '__main__':
    lang = sys.argv[1]
    passes = int(sys.argv[sys.argv.index('--passes') + 1]) if '--passes' in sys.argv else 2
    build(lang, passes)
